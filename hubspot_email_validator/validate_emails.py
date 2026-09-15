#!/usr/bin/env python3
"""
Email validator for HubSpot list exports.

Runs: RFC syntax check, live MX lookup, free-webmail detection,
role-address detection, and known typo-pattern detection (nifo@ etc).

Does NOT run an SMTP mailbox probe. For catch-all and dead-mailbox
detection, feed the SEND rows to a paid verifier.

Usage:
    pip install dnspython openpyxl pandas
    python validate_emails.py export.csv

Writes: <export>_validated.xlsx
"""

import re
import sys
import csv
import concurrent.futures as cf
from pathlib import Path

import dns.resolver
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SYNTAX = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")

FREE = {
    "gmail.com", "yahoo.com", "hotmail.com", "aol.com", "outlook.com",
    "icloud.com", "live.com", "msn.com", "ymail.com", "sbcglobal.net",
    "mail.com", "att.net", "me.com", "comcast.net", "verizon.net",
    "bellsouth.net", "earthlink.net", "protonmail.com", "gmx.com",
}

ROLE_PREFIX = (
    "info", "office", "sales", "admin", "service", "support", "contact",
    "accounting", "billing", "scheduling", "dispatch", "help", "hello",
    "accounts", "enquiries", "rentals", "noreply", "no-reply", "webmaster",
    "postmaster", "team", "careers", "hr", "jobs",
)

# Transposed / corrupted variants of info@ seen in the ZenTrades data
TYPO_LOCAL = {"nifo", "ifno", "inof", "ino", "onfo", "fino", "nfio"}

# Common domain typos worth catching
TYPO_DOMAIN = {
    "gmial.com", "gmai.com", "gmail.con", "gmail.co", "gnail.com",
    "yahoo.con", "yaho.com", "yahoo.co", "hotmial.com", "hotmai.com",
    "outlok.com", "iclould.com",
}


def find_email_column(header):
    for i, h in enumerate(header):
        if h and h.strip().lower() in ("email", "email address", "work email"):
            return i
    for i, h in enumerate(header):
        if h and "email" in h.strip().lower():
            return i
    raise SystemExit("No email column found. Expected a header containing 'email'.")


def load_emails(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        header = next(reader)
        idx = find_email_column(header)
        seen, out, dupes = set(), [], 0
        for row in reader:
            if idx >= len(row):
                continue
            e = row[idx].strip().lower()
            if not e:
                continue
            if e in seen:
                dupes += 1
                continue
            seen.add(e)
            out.append(e)
    print(f"Loaded {len(out)} unique addresses ({dupes} duplicates dropped)")
    return out


def make_resolver():
    r = dns.resolver.Resolver()
    r.lifetime = 8
    r.timeout = 4
    return r


def check_domain(domain):
    r = make_resolver()
    try:
        answers = r.resolve(domain, "MX")
        best = sorted(answers, key=lambda x: x.preference)[0]
        return domain, "MX_OK", str(best.exchange).rstrip(".")
    except dns.resolver.NXDOMAIN:
        return domain, "NXDOMAIN", ""
    except dns.resolver.NoAnswer:
        try:
            r.resolve(domain, "A")
            return domain, "NO_MX_HAS_A", ""
        except Exception:
            return domain, "NO_MX_NO_A", ""
    except Exception:
        return domain, "LOOKUP_TIMEOUT", ""


def classify(email, mxmap):
    local, _, domain = email.partition("@")
    status, host = mxmap.get(domain, ("LOOKUP_TIMEOUT", ""))
    syntax_ok = "YES" if SYNTAX.match(email) else "NO"
    free = "YES" if domain in FREE else ""
    role = "YES" if local.startswith(ROLE_PREFIX) else ""
    typo = "YES" if local in TYPO_LOCAL or domain in TYPO_DOMAIN else ""

    if syntax_ok == "NO" or status in ("NXDOMAIN", "NO_MX_NO_A") or typo == "YES":
        verdict = "SUPPRESS"
    elif status in ("NO_MX_HAS_A", "LOOKUP_TIMEOUT") or role == "YES":
        verdict = "REVIEW"
    else:
        verdict = "SEND"

    return [email, local, domain, syntax_ok, status, host, free, role, typo, verdict]


def write_workbook(rows, out_path):
    order = {"SEND": 0, "REVIEW": 1, "SUPPRESS": 2}
    rows.sort(key=lambda r: (order[r[9]], r[0]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validated"
    ws.append([
        "Email", "Local Part", "Domain", "Syntax Valid", "MX Status",
        "Primary MX Host", "Free Webmail", "Role Address", "Typo Flag", "Verdict",
    ])
    for r in rows:
        ws.append(r)

    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    body = Font(name="Arial")
    fills = {
        "SEND": PatternFill("solid", fgColor="D9EAD3"),
        "REVIEW": PatternFill("solid", fgColor="FFF2CC"),
        "SUPPRESS": PatternFill("solid", fgColor="F4CCCC"),
    }
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
        row[9].fill = fills[row[9].value]
        row[9].font = Font(name="Arial", bold=True)

    for i, w in enumerate([40, 22, 34, 13, 15, 34, 14, 14, 11, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"

    n = ws.max_row
    s = wb.create_sheet("Summary")
    s["A1"] = "Email validation summary"
    s["A1"].font = Font(name="Arial", bold=True, size=14)
    s["A2"] = "Checks: RFC syntax, live MX lookup, free webmail, role address, typo patterns"
    s["A3"] = "Not run: SMTP mailbox probe. Catch-all and dead-mailbox detection needs a paid verifier."
    for rr in s.iter_rows(min_row=2, max_row=3):
        for c in rr:
            c.font = Font(name="Arial", size=9)

    stats = [
        ("Rows checked", f"=COUNTA(Validated!A2:A{n})"),
        ("Verdict SEND", f'=COUNTIF(Validated!J2:J{n},"SEND")'),
        ("Verdict REVIEW", f'=COUNTIF(Validated!J2:J{n},"REVIEW")'),
        ("Verdict SUPPRESS", f'=COUNTIF(Validated!J2:J{n},"SUPPRESS")'),
        ("", ""),
        ("MX valid", f'=COUNTIF(Validated!E2:E{n},"MX_OK")'),
        ("No MX, A record only", f'=COUNTIF(Validated!E2:E{n},"NO_MX_HAS_A")'),
        ("Domain does not exist", f'=COUNTIF(Validated!E2:E{n},"NXDOMAIN")'),
        ("", ""),
        ("Free webmail", f'=COUNTIF(Validated!G2:G{n},"YES")'),
        ("Role addresses", f'=COUNTIF(Validated!H2:H{n},"YES")'),
        ("Typo addresses", f'=COUNTIF(Validated!I2:I{n},"YES")'),
    ]
    row_i = 5
    for label, formula in stats:
        s.cell(row_i, 1, label).font = Font(name="Arial", bold=bool(label))
        if formula:
            s.cell(row_i, 2, formula).font = Font(name="Arial")
        row_i += 1
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 14

    wb.save(out_path)


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python validate_emails.py export.csv")

    src = Path(sys.argv[1])
    emails = load_emails(src)

    domains = sorted({e.split("@")[1] for e in emails if "@" in e})
    print(f"Resolving MX for {len(domains)} unique domains...")

    mxmap = {}
    with cf.ThreadPoolExecutor(max_workers=40) as ex:
        for i, (d, status, host) in enumerate(ex.map(check_domain, domains), 1):
            mxmap[d] = (status, host)
            if i % 250 == 0:
                print(f"  {i}/{len(domains)}")

    rows = [classify(e, mxmap) for e in emails if "@" in e]

    out = src.with_name(src.stem + "_validated.xlsx")
    write_workbook(rows, out)

    counts = {}
    for r in rows:
        counts[r[9]] = counts.get(r[9], 0) + 1
    print("\nDone ->", out)
    for k in ("SEND", "REVIEW", "SUPPRESS"):
        print(f"  {k:9} {counts.get(k, 0)}")


if __name__ == "__main__":
    main()
